"""Create an Excel workbook summarising JSONL result files."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Tuple

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ResultSpec = Tuple[str, str]


def _load_results(file_path: Path) -> pd.DataFrame:
	"""Load a JSONL results file into a DataFrame with a sensible column order."""
	if not file_path.exists():
		raise FileNotFoundError(f"Missing results file: {file_path}")

	df = pd.read_json(file_path, lines=True)

	preferred_order = [col for col in ("prompt", "model", "output") if col in df.columns]
	remaining_columns = [col for col in df.columns if col not in preferred_order]

	if preferred_order:
		df = df[preferred_order + remaining_columns]

	return df


def _sanitize_table_name(name: str) -> str:
	"""Ensure the Excel table name starts with a letter and contains valid characters."""
	cleaned = "".join(ch for ch in name if ch.isalnum())
	if not cleaned:
		cleaned = "Results"
	if not cleaned[0].isalpha():
		cleaned = f"T{cleaned}"
	return cleaned[:253]


def _add_table_to_sheet(df: pd.DataFrame, sheet_name: str, workbook_writer: pd.ExcelWriter) -> None:
	"""Convert the written DataFrame range into an Excel table for styling and filtering."""
	worksheet = workbook_writer.sheets[sheet_name]

	row_count = len(df.index) + 1  # +1 for the header row
	col_count = len(df.columns)

	if col_count == 0:
		return  # Nothing to tabulate

	end_column = get_column_letter(col_count)
	table_ref = f"A1:{end_column}{row_count}"
	table_name = _sanitize_table_name(f"{sheet_name}Table")

	table = Table(displayName=table_name, ref=table_ref)
	table.tableStyleInfo = TableStyleInfo(
		name="TableStyleMedium9",
		showFirstColumn=False,
		showLastColumn=False,
		showRowStripes=True,
		showColumnStripes=False,
	)

	worksheet.add_table(table)


def build_results_workbook(
	output_path: Path,
	result_specs: Iterable[ResultSpec],
	base_dir: Path | None = None,
) -> Path:
	"""Create an Excel workbook with one sheet per JSONL results file."""
	base = base_dir or Path(__file__).parent

	output_path = output_path if output_path.is_absolute() else base / output_path
	output_path.parent.mkdir(parents=True, exist_ok=True)

	with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
		for filename, sheet_name in result_specs:
			df = _load_results(base / filename)
			df.to_excel(writer, sheet_name=sheet_name, index=False)
			_add_table_to_sheet(df, sheet_name, writer)

	return output_path


def main() -> None:
	"""Generate the combined results workbook in the project directory."""
	project_root = Path(__file__).parent
	workbook_path = build_results_workbook(
		output_path=Path("combined_results.xlsx"),
		result_specs=(
			("results_Balanced.jsonl", "Balanced"),
			("results_Cost.jsonl", "Cost"),
			("results_Quality.jsonl", "Quality"),
		),
		base_dir=project_root,
	)

	print(f"Excel workbook created at: {workbook_path}")


if __name__ == "__main__":
	main()
